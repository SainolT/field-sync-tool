# -*- coding: utf-8 -*-
"""字段映射同步工具 - 核心逻辑层（不含界面、不做 Excel 写入）

职责:
  · 工具表/配置的读写与校验
  · 表格扫描与映射有效性检查
  · 同步/替换的「变更计划」生成（纯匹配逻辑，只读数据文件）
  · 计划执行编排（备份 + 调用 xlsx_writer 外科手术写入）与统一日志模板

Excel 写入全部在 xlsx_writer.py（直接改 sheet XML，不动格式）。
"""
import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

import xlsx_writer

VERSION = "1.2.0"
TOOL_FILE = "字段同步工具表.xlsx"
MAIN_SHEET = "总表"
FIELD_SHEET = "字段表"
CONFIG_FILE = "工具配置.json"
BACKUP_DIR = "字段同步备份"
SCAN_CACHE_FILE = "扫描缓存.json"
MAIN_HEADERS = ["文件名称", "表格名称"]
FIELD_HEADERS = ["映射组", "文件名称", "表格名称", "字段名"]
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

INSTRUCTIONS = [
    ("字段映射同步工具 - 使用说明", True),
    ("", False),
    ("一、表格结构", True),
    ("「总表」: 登记表格清单, 只有两列 —— 文件名称 | 表格名称", False),
    ("「字段表」: 字段映射缓存(长表), 四列 —— 映射组 | 文件名称 | 表格名称 | 字段名", False),
    ("· 同一个映射组里的字段互为映射(如「作物名称」组里有 作物名称/茶树名称/茶树名)", False),
    ("· 按需缓存: 只登记你关心的字段, 没登记的字段不参与同步", False),
    ("· 数据文件按【工具表所在目录】查找, 支持任意层级的子目录, 文件名称显示相对路径(如 子目录/配置表.xlsx)", False),
    ("· 备份目录(默认「字段同步备份」)下的文件不会被扫描", False),
    ("· 含有「总表」「字段表」的文件会被识别为工具表, 不会当作数据表扫描", False),
    ("", False),
    ("二、同步操作", True),
    ("1. 打开工具(首次启动会引导你选择或新建一个表格作为工具表)", False),
    ("2. 选择「同步映射组」: 要同步的字段组(如 作物名称)", False),
    ("3. 选择「同步源字段」: 组内各字段已标明位置(文件|表格|字段名), 选一个作为数值标准", False),
    ("4. 在「被同步字段」里勾选目标(可点「全选」选中组内其余所有字段)", False),
    ("5. 选择「键映射组」: 用来对齐各表数据行的字段组(如 作物ID), 候选已按所选字段自动筛选", False),
    ("6. 点「预览」确认改动, 再点「开始同步」(默认自动备份)", False),
    ("· 备份目录默认为工具表旁的「字段同步备份」文件夹, 可点「设置备份目录」按钮更改, 选择会被记住", False),
    ("", False),
    ("三、登记字段(在工具窗口内完成)", True),
    ("1. 点「登记字段」按钮, 按 文件 → 子表 → 未登记字段 逐级选择(字段列表只显示还没登记的)", False),
    ("2. 勾选字段(可输入关键字筛选), 选择或输入一个映射组(可新建), 点「加入」", False),
    ("3. 一个表格可以分多次把不同字段加入不同映射组, 最后点「保存」写入工具表", False),
    ("· 也可以直接在「总表」「字段表」里手动加行, 效果相同; 注意同一表格在同一映射组里只能有一个字段, 重复添加视为修改(替换)", False),
    ("· 删除整个映射组: 点主窗口「删除映射组」按钮, 选组后确认(只删映射关系, 不动数据表)", False),
    ("", False),
    ("四、同步规则", True),
    ("· 按键组的值匹配行, 把数据源的值写入被同步表格在同一映射组里的字段", False),
    ("· 目标表格中所有键值相同的行都会更新", False),
    ("· 数据源中该值为空的行不同步(不会清空目标数据)", False),
    ("· 各表格的表头都应在第 1 行", False),
    ("· 所有操作都只改写数据区(第 2 行起), 不会修改任何表格的表头", False),
    ("· 公式保护: 源值是公式的不同步/不参照, 目标单元格是公式的不覆盖, 均会在日志给出警告", False),
    ("· 写入方式为外科手术式 XML 编辑: 只改目标单元格的值, 样式/格式/其它内容字节级不动", False),
    ("", False),
    ("五、字段替换(页签2)", True),
    ("共用: 「被替换字段」逐级选择 文件→子表→字段(所有字段都能选, 含未登记的; 可多选, 跨子表累积);", False),
    ("     输入关键字时切换为完整路径平铺筛选", False),
    ("方式一「按参照表(映射组)」: 声明 当前值是/替换成 两个映射组, 选一张同时登记了两组的表当字典;", False),
    ("  被替换列里的当前值去参照表查, 找到就替换。适合『ID列里临时填了名称, 批量翻回ID』", False),
    ("方式二「按字段匹配」: 选参照表格 + 匹配字段 + 取值字段; 按目标表同行的匹配字段值去参照表对齐行,", False),
    ("  把取值字段的值写入被替换列。目标子表需有与匹配字段同名的列", False),
    ("· 查不到对照的保留原样并列入日志; 空单元格跳过; 重复对照取第一行并警告", False),
    ("", False),
    ("六、常见问题", True),
    ("", False),
    ("Q1: 键ID字段含义相同、但在不同表里命名不同(如 作物ID / 茶树ID), 能用来匹配吗?", True),
    ("可以, 而且这就是映射组的设计用途。同步时每个表格用自己在键映射组里登记的字段名取键值,", False),
    ("再按键值匹配行, 字段名不同不影响。例如种植表的 作物ID=202001 和茶树表的 茶树ID=202001", False),
    ("会被认成同一行。", False),
    ("注意: 遇到还没归组的同义字段, 用「登记字段」把它加进【已有的】映射组, 不要新建组,", False),
    ("否则新旧字段就断开映射了。", False),
    ("", False),
    ("Q2: 键映射组误选了非ID类组(如 培育时间)会有什么影响?", True),
    ("不会报错, 但结果很可能是错的。键值不唯一时: 数据源里同键值只取第一行, 目标表里所有同键值", False),
    ("的行会被全部刷成那一行的值。判断标准: 该字段在每个参与表格里是否每行唯一(如 文本ID 就可以当键)。", False),
    ("防护: 同步前先点「预览」, 改动行数远超预期通常就是键选错了; 默认写入前会自动备份。", False),
]


# ============================== 配置与工具表 ==============================

def app_dir():
    """脚本/exe 自身所在目录（工具表和被管理的 xlsx 都放在这里）。"""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def load_config(base_dir):
    """读取工具配置（工具表路径、备份目录等），没有或损坏时返回空 dict。"""
    cfg = base_dir / CONFIG_FILE
    if not cfg.exists():
        return {}
    try:
        return json.loads(cfg.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config(base_dir, **kw):
    """更新工具配置（只覆盖传入的键）。"""
    data = load_config(base_dir)
    data.update(kw)
    (base_dir / CONFIG_FILE).write_text(
        json.dumps(data, ensure_ascii=False), encoding="utf-8")


def _config_path(base_dir, value):
    """配置里的路径: 相对路径相对 app 目录解析。"""
    p = Path(value)
    if not p.is_absolute():
        p = base_dir / p
    return p


def find_tool_config(base_dir):
    """读取上次选择的工具表路径，没有配置或文件已不存在时返回 None。"""
    v = load_config(base_dir).get("tool_file", "")
    if not v:
        return None
    p = _config_path(base_dir, v)
    return p if p.exists() else None


def save_tool_config(base_dir, path):
    try:
        s = str(path.relative_to(base_dir))
    except ValueError:
        s = str(path)
    save_config(base_dir, tool_file=s)


def find_backup_config(base_dir):
    """读取用户设置的备份目录，未设置时返回 None。"""
    v = load_config(base_dir).get("backup_dir", "")
    return _config_path(base_dir, v) if v else None


def save_backup_config(base_dir, path):
    try:
        s = str(path.relative_to(base_dir))
    except ValueError:
        s = str(path)
    save_config(base_dir, backup_dir=s)


def norm(v):
    """把单元格值归一化成可比较的字符串（1 / 1.0 / '1' 视为相同）。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def is_formula_like(v):
    """判断值是否形如公式（以 = 开头的文本，写入后会变成公式，有注入风险）。"""
    return isinstance(v, str) and v.startswith("=")


def is_formula_cell(cell):
    """判断单元格是否是公式: 覆盖普通公式、数组公式(ArrayFormula 对象)和以 = 开头的文本。"""
    return cell.data_type == "f" or is_formula_like(cell.value)


def header_map_of(ws):
    """读取已打开工作表的第 1 行表头，返回 {列名: 列号(1-based)}。"""
    hmap = {}
    for c in ws[1]:
        h = norm(c.value)
        if h and h not in hmap:
            hmap[h] = c.column
    return hmap


def read_sheet(path, sheet):
    """打开目标表格，返回 (workbook, worksheet, {列名: 列号})。"""
    wb = openpyxl.load_workbook(path)
    if sheet not in wb.sheetnames:
        raise KeyError(f"工作表「{sheet}」不存在于 {Path(path).name}")
    ws = wb[sheet]
    return wb, ws, header_map_of(ws)


def check_tool_format(path):
    """检查表格: 'ok'=符合工具表格式, 'empty'=空表(可初始化), 'invalid'=两者都不是。"""
    try:
        wb = openpyxl.load_workbook(path)
    except Exception:
        return "invalid"
    try:
        has_data = any(
            any(c.value is not None for c in row)
            for ws in wb.worksheets
            for row in ws.iter_rows(max_row=1)
        ) or any(ws.max_row > 1 for ws in wb.worksheets)
        if not has_data:
            return "empty"
        if MAIN_SHEET in wb.sheetnames and FIELD_SHEET in wb.sheetnames:
            h1 = [norm(c.value) for c in wb[MAIN_SHEET][1][:2]]
            h2 = [norm(c.value) for c in wb[FIELD_SHEET][1][:4]]
            if h1 == MAIN_HEADERS and h2 == FIELD_HEADERS:
                return "ok"
        return "invalid"
    finally:
        wb.close()


def init_tool_file(path):
    """把表格初始化为工具表: 写入使用说明, 确保「总表」「字段表」及其表头存在。"""
    path = Path(path)
    if path.exists():
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "使用说明"
    if "使用说明" in wb.sheetnames:
        del wb["使用说明"]
    ws = wb.create_sheet("使用说明", 0)
    for i, (text, bold) in enumerate(INSTRUCTIONS, 1):
        c = ws.cell(i, 1, text)
        if bold:
            c.font = Font(bold=True, size=12 if "、" not in text else 11)
    ws.column_dimensions["A"].width = 95
    for name, headers, widths in [
        (MAIN_SHEET, MAIN_HEADERS, [26, 24]),
        (FIELD_SHEET, FIELD_HEADERS, [16, 26, 24, 18]),
    ]:
        if name in wb.sheetnames:
            ws = wb[name]
        else:
            ws = wb.create_sheet(name)
        for j, h in enumerate(headers, 1):
            c = ws.cell(1, j, h)
            c.font = Font(bold=True)
            c.fill = HEADER_FILL
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    wb.save(path)


def load_tool(path):
    """读取工具表，返回 (tables, groups)。
    tables: [(文件, 表格), ...]（保持总表顺序）
    groups: {映射组: {(文件, 表格): 字段名}}"""
    wb = openpyxl.load_workbook(path, read_only=True)
    if MAIN_SHEET not in wb.sheetnames or FIELD_SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(f"工具表需要包含「{MAIN_SHEET}」和「{FIELD_SHEET}」两个工作表")
    tables = []
    for r in wb[MAIN_SHEET].iter_rows(min_row=2, values_only=True):
        if r and norm(r[0] if len(r) > 0 else None) and norm(r[1] if len(r) > 1 else None):
            tables.append((norm(r[0]), norm(r[1])))
    groups = {}
    for r in wb[FIELD_SHEET].iter_rows(min_row=2, values_only=True):
        if not r or len(r) < 4:
            continue
        g, f, s, col = (norm(r[0]), norm(r[1]), norm(r[2]), norm(r[3]))
        if g and f and s and col:
            groups.setdefault(g, {})[(f, s)] = col
    wb.close()
    return tables, groups


# ============================== 扫描与映射管理 ==============================

def valid_headers(ws):
    hdr = list(next(ws.iter_rows(max_row=1, values_only=True), None) or [])
    while hdr and hdr[-1] is None:
        hdr.pop()
    if not hdr or not all(isinstance(h, str) and h.strip() for h in hdr):
        return None
    if len(set(hdr)) != len(hdr):
        return None
    return [h.strip() for h in hdr]


def scan_tables(base_dir, tool_path=None, exclude_dir=None, cache_dir=None):
    """扫描根目录（含任意层级子目录）下所有可用的 xlsx 表格。
    返回 [(相对路径, [(子表名, [表头...]), ...]), ...]，按文件分组。
    exclude_dir（备份目录）下的文件不参与。
    cache_dir 提供时启用扫描缓存（按 路径+修改时间+大小 判断变化），
    未变化的文件不重复解析，大幅加快多文件场景的扫描。"""
    result = []
    base_dir = Path(base_dir)
    tool_resolved = Path(tool_path).resolve() if tool_path else None
    excl_resolved = Path(exclude_dir).resolve() if exclude_dir else None
    cache = {}
    cache_path = Path(cache_dir) / SCAN_CACHE_FILE if cache_dir else None
    if cache_path and cache_path.exists():
        try:
            cache = json.loads(cache_path.read_text(encoding="utf-8"))
        except Exception:
            cache = {}
    new_cache = {}
    for p in sorted(base_dir.rglob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        rp = p.resolve()
        if tool_resolved and rp == tool_resolved:
            continue
        if tool_resolved is None and p.name == TOOL_FILE:
            continue
        if excl_resolved and excl_resolved in rp.parents:
            continue
        try:
            st = p.stat()
        except OSError:
            continue
        key = str(rp)
        ent = cache.get(key)
        if ent and ent.get("mtime") == st.st_mtime and ent.get("size") == st.st_size:
            new_cache[key] = ent
        else:
            ent = {"mtime": st.st_mtime, "size": st.st_size, "tool": False, "sheets": []}
            try:
                wb = openpyxl.load_workbook(p, read_only=True)
            except Exception:
                new_cache[key] = ent
                continue
            if MAIN_SHEET in wb.sheetnames and FIELD_SHEET in wb.sheetnames:
                ent["tool"] = True  # 工具表格式的文件不作为数据表
            else:
                for ws in wb.worksheets:
                    hdr = valid_headers(ws)
                    if hdr:
                        ent["sheets"].append([ws.title, hdr])
            wb.close()
            new_cache[key] = ent
        if ent["tool"] or not ent["sheets"]:
            continue
        result.append((p.relative_to(base_dir).as_posix(), ent["sheets"]))
    if cache_path:
        try:
            cache_path.write_text(json.dumps(new_cache, ensure_ascii=False), encoding="utf-8")
        except Exception:
            pass
    return result


def check_mappings(base_dir, tables, groups, cache_dir=None, tool_path=None, exclude_dir=None):
    """检查字段表里的映射是否仍然有效（文件/工作表/字段列都存在）。
    返回失效映射列表 [(映射组, 文件, 表格, 字段名, 原因), ...]。
    表头信息走 scan_tables 的扫描缓存，多文件场景下更快。"""
    lost = []
    registered = set(tables)
    scanned = {
        f: {s: hdr for s, hdr in sheets}
        for f, sheets in scan_tables(base_dir, tool_path, exclude_dir, cache_dir)
    }
    for g, mapping in groups.items():
        for (f, s), col in mapping.items():
            if (f, s) not in registered:
                lost.append((g, f, s, col, "表格未在总表登记"))
                continue
            if not (base_dir / f).exists():
                lost.append((g, f, s, col, "文件不存在"))
                continue
            info = scanned.get(f)
            if not info:
                lost.append((g, f, s, col, "文件无法读取"))
            elif s not in info:
                lost.append((g, f, s, col, "工作表不存在"))
            elif col not in info[s]:
                lost.append((g, f, s, col, "字段列不存在"))
    return lost


def add_associations(tool_path, additions):
    """把新登记写入工具表。additions: [(文件, 表格, 字段名, 映射组), ...]。
    同一映射组内同一表格只保留一个字段：重复添加视为修改（替换字段名）。
    返回 (新增行数, 更新行数)。"""
    wb = openpyxl.load_workbook(tool_path)
    ws_main, ws_field = wb[MAIN_SHEET], wb[FIELD_SHEET]
    registered = set()
    for r in ws_main.iter_rows(min_row=2, values_only=True):
        if r and norm(r[0] if len(r) > 0 else None) and norm(r[1] if len(r) > 1 else None):
            registered.add((norm(r[0]), norm(r[1])))
    existing = {}  # (映射组, 文件, 表格) -> 行号
    for i, r in enumerate(ws_field.iter_rows(min_row=2, values_only=True), start=2):
        if r and len(r) >= 4 and norm(r[0]) and norm(r[1]) and norm(r[2]):
            existing[(norm(r[0]), norm(r[1]), norm(r[2]))] = i
    added = updated = 0
    for f, s, col, g in additions:
        if (f, s) not in registered:
            ws_main.append([f, s])
            registered.add((f, s))
        key = (g, f, s)
        if key in existing:
            row_i = existing[key]
            if norm(ws_field.cell(row_i, 4).value) != col:
                ws_field.cell(row_i, 4).value = col
                updated += 1
        else:
            ws_field.append([g, f, s, col])
            existing[key] = ws_field.max_row
            added += 1
    wb.save(tool_path)
    return added, updated


def remove_mappings(tool_path, lost):
    """从字段表删除失效映射行，返回删除条数。"""
    targets = {(g, f, s, col) for g, f, s, col, _ in lost}
    wb = openpyxl.load_workbook(tool_path)
    ws = wb[FIELD_SHEET]
    n = 0
    for r in range(ws.max_row, 1, -1):
        vals = tuple(norm(ws.cell(r, c).value) for c in range(1, 5))
        if vals in targets:
            ws.delete_rows(r)
            n += 1
    wb.save(tool_path)
    return n


def delete_group(tool_path, group):
    """删除整个映射组（字段表中该组的所有行），返回删除条数。"""
    wb = openpyxl.load_workbook(tool_path)
    ws = wb[FIELD_SHEET]
    n = 0
    for r in range(ws.max_row, 1, -1):
        if norm(ws.cell(r, 1).value) == group:
            ws.delete_rows(r)
            n += 1
    wb.save(tool_path)
    return n


# ============================== 变更计划与执行 ==============================

def _new_plan():
    return {"fatal": None, "detail": [], "warn": [], "files": {}, "open_failed": [],
            "planned": 0, "blocked": 0, "unmatched": 0}


def _logger(detail, warn):
    def log(msg):
        (warn if msg.strip().startswith("[警告]") else detail).append(msg)
    return log


def plan_sync(base_dir, groups, key_group, src, sync_groups, targets):
    """生成同步计划(只读, 不写盘)。
    files: {文件: {子表: [(行, 列号, 旧值, 新值), ...]}}"""
    res = _new_plan()
    log = _logger(res["detail"], res["warn"])
    src_key_col = groups.get(key_group, {}).get(src)
    if not src_key_col:
        res["fatal"] = [f"[错误] 数据源表格在映射组「{key_group}」里没有登记字段"]
        return res
    try:
        swb, sws, shdr = read_sheet(base_dir / src[0], src[1])
    except Exception as ex:
        res["fatal"] = [f"[错误] 读取数据源失败: {ex}"]
        return res
    if src_key_col not in shdr:
        res["fatal"] = [f"[错误] 数据源表头里找不到键列「{src_key_col}」"]
        return res

    src_fields = {}
    for g in sync_groups:
        col = groups.get(g, {}).get(src)
        if not col:
            log(f"[跳过] 数据源表格在映射组「{g}」里没有登记字段")
            continue
        if col not in shdr:
            log(f"[跳过] 数据源表头里找不到字段列「{col}」(映射组「{g}」)")
            continue
        src_fields[g] = shdr[col]
    if not src_fields:
        res["fatal"] = ["[中止] 数据源没有任何可同步的字段"]
        return res

    key_c = shdr[src_key_col]
    src_data = {}
    dup_keys = 0
    formula_src = []
    for r in range(2, sws.max_row + 1):
        k = norm(sws.cell(r, key_c).value)
        if not k:
            continue
        if k in src_data:
            dup_keys += 1
            continue
        vals = {}
        for g, c in src_fields.items():
            cell = sws.cell(r, c)
            v = cell.value
            if is_formula_cell(cell):
                formula_src.append((r, g))
                v = None  # 公式不同步，视同空值跳过
            vals[g] = v
        src_data[k] = vals
    swb.close()
    res["blocked"] += len(formula_src)
    if formula_src:
        sample = ", ".join(f"行{r}「{g}」" for r, g in formula_src[:5])
        log(f"[警告] 数据源有 {len(formula_src)} 个值是公式，已跳过（工具不同步公式内容）: {sample}"
            + (" ……" if len(formula_src) > 5 else ""))
    if dup_keys:
        log(f"[提示] 数据源有 {dup_keys} 个重复的键值，已取首次出现的行")
    log(f"数据源: {src[0]} | {src[1]}，共 {len(src_data)} 个键值")
    log(f"键映射组: {key_group} | 同步映射组: {', '.join(src_fields)}")

    by_file = {}
    for tgt in targets:
        if tgt == src:
            continue
        by_file.setdefault(tgt[0], []).append(tgt)

    for f, tgts in by_file.items():
        res["files"][f] = {}
        try:
            twb = openpyxl.load_workbook(base_dir / f)
        except Exception as ex:
            log(f"[错误] {f}: 打开失败: {ex}")
            res["open_failed"].append(f)
            continue
        for tgt in tgts:
            s = tgt[1]
            label = f"{f} | {s}"
            tgt_key_col = groups.get(key_group, {}).get(tgt)
            if not tgt_key_col:
                log(f"[跳过] {label}: 在映射组「{key_group}」里没有登记字段")
                continue
            if s not in twb.sheetnames:
                log(f"[错误] {label}: 工作表不存在")
                continue
            tws = twb[s]
            thdr = header_map_of(tws)
            if tgt_key_col not in thdr:
                log(f"[跳过] {label}: 表头里找不到键列「{tgt_key_col}」")
                continue
            tgt_fields = {}
            for g in src_fields:
                col = groups.get(g, {}).get(tgt)
                if not col:
                    log(f"[提示] {label}: 在映射组「{g}」里没有登记字段，该组跳过")
                    continue
                if col not in thdr:
                    log(f"[提示] {label}: 表头里找不到字段列「{col}」，该组跳过")
                    continue
                tgt_fields[g] = thdr[col]
            if not tgt_fields:
                log(f"[跳过] {label}: 没有可同步的字段")
                continue

            kc = thdr[tgt_key_col]
            changes = []
            hit_keys = set()
            formula_tgt = 0
            for r in range(2, tws.max_row + 1):
                k = norm(tws.cell(r, kc).value)
                if not k or k not in src_data:
                    continue
                hit_keys.add(k)
                for g, c in tgt_fields.items():
                    v = src_data[k][g]
                    if v is None or norm(v) == "":
                        continue  # 源值为空不同步
                    cell = tws.cell(r, c)
                    if is_formula_cell(cell):
                        formula_tgt += 1  # 目标是公式，不覆盖
                        continue
                    if norm(cell.value) != norm(v):
                        changes.append((r, g, cell.value, v))
            res["blocked"] += formula_tgt
            res["files"][f][s] = [(r, tgt_fields[g], old, new) for r, g, old, new in changes]
            res["planned"] += len(changes)
            if changes:
                log(f"\n{label}: {len(changes)} 处改动")
                for r, g, old, new in changes[:200]:
                    log(f"  行{r} 「{groups[g][tgt]}」: {old!r} -> {new!r}")
                if len(changes) > 200:
                    log(f"  …… 其余 {len(changes) - 200} 处省略")
            else:
                log(f"{label}: 无需改动（值已一致）")
            if formula_tgt:
                log(f"  [警告] {label}: {formula_tgt} 个目标单元格是公式，未覆盖（工具不改动公式）")
            missing = set(src_data) - hit_keys
            if missing:
                res["unmatched"] += len(missing)
                sample = ", ".join(sorted(missing)[:10])
                log(f"  [提示] {label}: 有 {len(missing)} 个数据源的键值未匹配到行（{sample}{' ……' if len(missing) > 10 else ''}）")
        twb.close()
    return res


def plan_replace(base_dir, groups, src_group, dst_group, ref, targets):
    """按参照表替换的计划(只读, 不写盘)。targets: [(文件, 表格, 字段名), ...]"""
    res = _new_plan()
    log = _logger(res["detail"], res["warn"])
    src_col = groups.get(src_group, {}).get(ref)
    dst_col = groups.get(dst_group, {}).get(ref)
    if not src_col or not dst_col:
        res["fatal"] = ["[错误] 参照表没有同时登记源组和目标组的字段"]
        return res
    try:
        rwb, rws, rhdr = read_sheet(base_dir / ref[0], ref[1])
    except Exception as ex:
        res["fatal"] = [f"[错误] 读取参照表失败: {ex}"]
        return res
    if src_col not in rhdr or dst_col not in rhdr:
        res["fatal"] = [f"[错误] 参照表表头里找不到列「{src_col}」或「{dst_col}」"]
        rwb.close()
        return res
    mapping = {}
    dup = 0
    formula_ref = 0
    for r in range(2, rws.max_row + 1):
        k = norm(rws.cell(r, rhdr[src_col]).value)
        if not k:
            continue
        if k in mapping:
            dup += 1
            continue
        cell = rws.cell(r, rhdr[dst_col])
        v_dst = cell.value
        if is_formula_cell(cell):
            formula_ref += 1
            v_dst = None  # 公式不作为参照值，视同空值跳过
        mapping[k] = v_dst
    dst_vals = {norm(v) for v in mapping.values() if norm(v)}
    rwb.close()
    res["blocked"] += formula_ref
    if dup:
        log(f"[提示] 参照表有 {dup} 个重复的「{src_group}」值，已取首次出现的行")
    if formula_ref:
        log(f"[警告] 参照表有 {formula_ref} 条对照的取值是公式，已跳过（工具不处理公式内容）")
    log(f"参照表: {ref[0]} | {ref[1]}（{src_col} → {dst_col}），共 {len(mapping)} 条对照")
    log(f"翻译规则: 「{src_group}」→「{dst_group}」")

    by_file = {}
    for f, s, col in targets:
        by_file.setdefault(f, []).append((s, col))

    for f, entries in by_file.items():
        res["files"][f] = {}
        try:
            twb = openpyxl.load_workbook(base_dir / f)
        except Exception as ex:
            log(f"[错误] {f}: 打开失败: {ex}")
            res["open_failed"].append(f)
            continue
        for s, col in entries:
            label = f"{f} | {s} | {col}"
            if s not in twb.sheetnames:
                log(f"[错误] {label}: 工作表不存在")
                continue
            tws = twb[s]
            thdr = header_map_of(tws)
            if col not in thdr:
                log(f"[跳过] {label}: 表头里找不到该列")
                continue
            c = thdr[col]
            changes = []
            unmatched = {}
            already = 0
            empty_dst = 0
            formula_tgt = 0
            for r in range(2, tws.max_row + 1):
                cell = tws.cell(r, c)
                v = cell.value
                if v is None or norm(v) == "":
                    continue
                if is_formula_cell(cell):
                    formula_tgt += 1  # 目标是公式(含数组公式)，不替换
                    continue
                k = norm(v)
                if k in mapping:
                    pass  # 待替换
                elif k in dst_vals:
                    already += 1  # 已是目标值，跳过
                    continue
                else:
                    unmatched[k] = unmatched.get(k, 0) + 1
                    continue
                new = mapping[k]
                if new is None or norm(new) == "":
                    empty_dst += 1
                    continue
                if norm(v) != norm(new):
                    changes.append((r, v, new))
            res["blocked"] += formula_tgt
            res["files"][f][f"{s}::{col}"] = [(r, c, old, new) for r, old, new in changes]
            res["planned"] += len(changes)
            if changes:
                log(f"\n{label}: {len(changes)} 处替换")
                for r, old, new in changes[:200]:
                    log(f"  行{r}: {old!r} -> {new!r}")
                if len(changes) > 200:
                    log(f"  …… 其余 {len(changes) - 200} 处省略")
            else:
                log(f"{label}: 无需替换")
            if unmatched:
                res["unmatched"] += sum(unmatched.values())
                items = ", ".join(f"{k!r}×{n}" for k, n in sorted(unmatched.items())[:10])
                log(f"  [提示] {label}: {sum(unmatched.values())} 个值在参照表找不到（保留原值）: {items}"
                    + (" ……" if len(unmatched) > 10 else ""))
            if already:
                log(f"  [提示] {label}: {already} 个单元格已是「{dst_group}」的值，已跳过")
            if formula_tgt:
                log(f"  [警告] {label}: {formula_tgt} 个目标单元格是公式，未替换（工具不改动公式）")
            if empty_dst:
                log(f"  [提示] {label}: {empty_dst} 个值在参照表中的目标列为空（保留原值）")
        twb.close()
    return res


def plan_replace_align(base_dir, ref, match_col, value_col, targets):
    """按字段匹配替换的计划(只读, 不写盘)。targets: [(文件, 表格, 字段名), ...]"""
    res = _new_plan()
    log = _logger(res["detail"], res["warn"])
    try:
        rwb, rws, rhdr = read_sheet(base_dir / ref[0], ref[1])
    except Exception as ex:
        res["fatal"] = [f"[错误] 读取参照表失败: {ex}"]
        return res
    if match_col not in rhdr or value_col not in rhdr:
        res["fatal"] = [f"[错误] 参照表表头里找不到列「{match_col}」或「{value_col}」"]
        rwb.close()
        return res
    mapping = {}
    dup = 0
    formula_ref = 0
    for r in range(2, rws.max_row + 1):
        k = norm(rws.cell(r, rhdr[match_col]).value)
        if not k:
            continue
        if k in mapping:
            dup += 1
            continue
        cell = rws.cell(r, rhdr[value_col])
        v_dst = cell.value
        if is_formula_cell(cell):
            formula_ref += 1
            v_dst = None  # 公式不作为参照值，视同空值跳过
        mapping[k] = v_dst
    rwb.close()
    res["blocked"] += formula_ref
    if dup:
        log(f"[提示] 参照表有 {dup} 个重复的「{match_col}」值，已取首次出现的行")
    if formula_ref:
        log(f"[警告] 参照表有 {formula_ref} 条对照的取值是公式，已跳过（工具不处理公式内容）")
    log(f"参照表: {ref[0]} | {ref[1]}（按「{match_col}」对齐，取「{value_col}」），共 {len(mapping)} 条对照")

    by_file = {}
    for f, s, col in targets:
        by_file.setdefault(f, []).append((s, col))

    for f, entries in by_file.items():
        res["files"][f] = {}
        try:
            twb = openpyxl.load_workbook(base_dir / f)
        except Exception as ex:
            log(f"[错误] {f}: 打开失败: {ex}")
            res["open_failed"].append(f)
            continue
        for s, col in entries:
            label = f"{f} | {s} | {col}"
            if s not in twb.sheetnames:
                log(f"[错误] {label}: 工作表不存在")
                continue
            tws = twb[s]
            thdr = header_map_of(tws)
            if col not in thdr:
                log(f"[跳过] {label}: 表头里找不到该列")
                continue
            if match_col not in thdr:
                log(f"[跳过] {label}: 该表没有匹配列「{match_col}」，无法对齐")
                continue
            if col == match_col:
                log(f"[跳过] {label}: 目标列与匹配列相同")
                continue
            mc, c = thdr[match_col], thdr[col]
            changes = []
            unmatched = {}
            already = 0
            empty_dst = 0
            formula_tgt = 0
            for r in range(2, tws.max_row + 1):
                kv = norm(tws.cell(r, mc).value)
                if not kv:
                    continue
                if kv not in mapping:
                    unmatched[kv] = unmatched.get(kv, 0) + 1
                    continue
                new = mapping[kv]
                if new is None or norm(new) == "":
                    empty_dst += 1
                    continue
                cell = tws.cell(r, c)
                if is_formula_cell(cell):
                    formula_tgt += 1  # 目标是公式(含数组公式)，不替换
                    continue
                if norm(cell.value) == norm(new):
                    already += 1
                    continue
                changes.append((r, cell.value, new))
            res["blocked"] += formula_tgt
            res["files"][f][f"{s}::{col}"] = [(r, c, old, new) for r, old, new in changes]
            res["planned"] += len(changes)
            if changes:
                log(f"\n{label}: {len(changes)} 处替换")
                for r, old, new in changes[:200]:
                    log(f"  行{r}: {old!r} -> {new!r}")
                if len(changes) > 200:
                    log(f"  …… 其余 {len(changes) - 200} 处省略")
            else:
                log(f"{label}: 无需替换")
            if unmatched:
                res["unmatched"] += sum(unmatched.values())
                items = ", ".join(f"{k!r}×{n}" for k, n in sorted(unmatched.items())[:10])
                log(f"  [提示] {label}: {sum(unmatched.values())} 行的「{match_col}」在参照表找不到（保留原值）: {items}"
                    + (" ……" if len(unmatched) > 10 else ""))
            if already:
                log(f"  [提示] {label}: {already} 个单元格已与参照值一致，已跳过")
            if formula_tgt:
                log(f"  [警告] {label}: {formula_tgt} 个目标单元格是公式，未替换（工具不改动公式）")
            if empty_dst:
                log(f"  [提示] {label}: {empty_dst} 行在参照表中的「{value_col}」为空（保留原值）")
        twb.close()
    return res


def emit_report(log, op, config_line, detail, warn, results, planned, saved,
                blocked, unmatched, failed, dry, unit="改动"):
    """统一日志模板输出。"""
    out = [f"【{op}】{datetime.now():%Y-%m-%d %H:%M:%S}", f"配置: {config_line}", "─" * 46, "明细:"]
    out.extend(f"  {l}" for l in detail)
    out.append("警告:")
    if warn:
        out.extend(f"  {l}" for l in warn)
    else:
        out.append("  无")
    out.append("─" * 46)
    out.append("执行结果:")
    out.extend(f"  {l}" for l in results)
    if dry:
        status = "预览(未写入)"
    elif failed == 0:
        status = "成功"
    elif saved == 0:
        status = "失败"
    else:
        status = "部分失败"
    out.append(f"汇总: 计划{unit} {planned} | 实际保存 {saved} | 公式拦截 {blocked} | "
               f"未匹配 {unmatched} | 失败文件 {failed} | 状态: {status}")
    log("\n".join(out))


def _apply_plan(base_dir, plan, backup_dir, dry, detail, warn, unit="改动"):
    """执行计划: 每文件备份一次 + XML 外科手术写入。返回 (results, saved, failed)。"""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    results, saved, failed = [], 0, 0
    for f in plan["open_failed"]:
        failed += 1
        results.append(f"{f}: 打开失败，未处理")
    for f, sheets in plan["files"].items():
        if f in plan["open_failed"]:
            continue
        planned_f = sum(len(v) for v in sheets.values())
        if not planned_f:
            results.append(f"{f}: 无需改动")
            continue
        if dry:
            results.append(f"{f}: 预览 {planned_f} 处{unit}")
            continue
        # sheets 的键可能是 "子表::字段"（替换）或 "子表"（同步），都按子表聚合;
        # 计划条目是 (行, 列号, 旧值, 新值), 写入器只需要 (行, 列号, 新值)
        by_sheet = {}
        for key, changes in sheets.items():
            sheet = key.split("::")[0]
            by_sheet.setdefault(sheet, []).extend((r, c, new) for r, c, old, new in changes)
        try:
            bak, applied, missed = xlsx_writer.save_file_with_changes(
                base_dir, f, by_sheet, backup_dir, stamp)
            saved += applied
            detail.append(f"[已保存] {f}（本批共 {applied} 处{unit}，已按原始状态备份）")
            results.append(f"{f}: {applied} 处{unit}，已保存"
                           + (f"（备份: {bak}）" if bak else "（未备份）"))
            for m in missed:
                warn.append(f"[警告] {f}: 未命中 {m}")
        except Exception as ex:
            failed += 1
            detail.append(f"[错误] {f}: 保存失败: {ex}（原文件未被修改）")
            results.append(f"{f}: 保存失败（{ex}），原文件未被修改")
    return results, saved, failed


def run_sync(base_dir, groups, key_group, src, sync_groups, targets, backup_dir, dry, log):
    """同步: 计划 → (备份+XML写入) → 统一模板报告。返回计划改动单元格数。"""
    p = plan_sync(base_dir, groups, key_group, src, sync_groups, targets)
    if p["fatal"]:
        for l in p["fatal"]:
            log(l)
        return 0
    results, saved, failed = _apply_plan(base_dir, p, backup_dir, dry, p["detail"], p["warn"])
    emit_report(log, "同步",
                f"键={key_group} | 源={src[0]}|{src[1]} | 同步组={', '.join(sync_groups)} | 目标={len(p['files'])} 个文件",
                p["detail"], p["warn"], results,
                p["planned"], saved, p["blocked"], p["unmatched"], failed, dry)
    return p["planned"]


def run_replace(base_dir, groups, src_group, dst_group, ref, targets, backup_dir, dry, log):
    """按参照表替换: 计划 → (备份+XML写入) → 统一模板报告。返回计划替换单元格数。"""
    p = plan_replace(base_dir, groups, src_group, dst_group, ref, targets)
    if p["fatal"]:
        for l in p["fatal"]:
            log(l)
        return 0
    results, saved, failed = _apply_plan(base_dir, p, backup_dir, dry, p["detail"], p["warn"], unit="替换")
    emit_report(log, "替换-按参照表",
                f"规则={src_group}→{dst_group} | 参照={ref[0]}|{ref[1]} | 目标={len(p['files'])} 个文件",
                p["detail"], p["warn"], results,
                p["planned"], saved, p["blocked"], p["unmatched"], failed, dry, unit="替换")
    return p["planned"]


def run_replace_align(base_dir, ref, match_col, value_col, targets, backup_dir, dry, log):
    """按字段匹配替换: 计划 → (备份+XML写入) → 统一模板报告。返回计划替换单元格数。"""
    p = plan_replace_align(base_dir, ref, match_col, value_col, targets)
    if p["fatal"]:
        for l in p["fatal"]:
            log(l)
        return 0
    results, saved, failed = _apply_plan(base_dir, p, backup_dir, dry, p["detail"], p["warn"], unit="替换")
    emit_report(log, "替换-按字段匹配",
                f"参照={ref[0]}|{ref[1]} | 对齐={match_col} | 取值={value_col} | 目标={len(p['files'])} 个文件",
                p["detail"], p["warn"], results,
                p["planned"], saved, p["blocked"], p["unmatched"], failed, dry, unit="替换")
    return p["planned"]
